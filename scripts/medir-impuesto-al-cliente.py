# -*- coding: utf-8 -*-
"""Que impuesto se le cobro de verdad al cliente en cada cotizacion.

Existe porque el 13% no siempre aplica. La hoja `Regimen Fiscal Clientes`
del catalogo tiene cuatro entradas --CCSS exenta, UCR y TEC al 2%, zona
franca variable-- y la pregunta es si con eso alcanza.

En vez de tratar de acordarse de que entidades son especiales, esto lo saca
de las cotizaciones: en la pestana que ve el cliente estan el SUBTOTAL y el
IMPUESTO, asi que la tasa efectiva es una division. Contando cuantas veces
aparece cada tasa se ve que regimenes existen de verdad y cuales nos faltan.

Solo lee.

Uso:
    python scripts/medir-impuesto-al-cliente.py
"""

import collections
import importlib.util
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")

_RUTA_BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "medir-base-de-costo.py")
_spec = importlib.util.spec_from_file_location("medir_base", _RUTA_BASE)
_mb = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mb)

RX_SUB = re.compile(r"^sub\s*-?\s*total", re.I)
RX_IMP = re.compile(r"^(impuesto|iva|i\.v\.a)", re.I)


def totales_de(filas):
    """Busca SUBTOTAL e IMPUESTO en una hoja y devuelve sus montos.

    El monto no esta en la celda de la etiqueta sino a su derecha, y no
    siempre en la columna de al lado: las plantillas varian. Se toma el
    primer numero que aparezca a la derecha, en la misma fila.
    """
    sub = imp = None
    for fila in filas:
        for c, v in enumerate(fila):
            if not isinstance(v, str):
                continue
            t = v.strip()
            objetivo = None
            if sub is None and RX_SUB.match(t):
                objetivo = "sub"
            elif imp is None and RX_IMP.match(t) and len(t) < 20:
                objetivo = "imp"
            if not objetivo:
                continue
            for w in fila[c + 1:]:
                if isinstance(w, (int, float)) and w != 0:
                    if objetivo == "sub":
                        sub = float(w)
                    else:
                        imp = float(w)
                    break
    return sub, imp


def main():
    archivos = []
    for root, dirs, files in os.walk(_mb.CLIENTES):
        if "00_IA_PREVENTAS" in root:
            continue
        for f in files:
            if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not re.search(r"matriz|oferta", f, re.I):
                continue
            p = os.path.join(root, f)
            try:
                if os.path.getsize(p) < 40000:
                    continue
            except OSError:
                continue
            archivos.append(p)
    print("matrices a revisar: %d" % len(archivos))

    tasas = collections.Counter()
    cta = collections.Counter()
    raros = []
    for i, p in enumerate(archivos, start=1):
        if i % 150 == 0:
            print("   %d/%d   tasas leidas: %d" % (i, len(archivos), sum(tasas.values())))
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception:
            cta["ilegibles"] += 1
            continue
        visto = False
        try:
            for nombre in wb.sheetnames:
                if not re.search(r"cotiza", nombre, re.I):
                    continue
                try:
                    filas = list(wb[nombre].iter_rows(min_row=1, max_row=90,
                                                      values_only=True))
                except Exception:
                    continue
                sub, imp = totales_de(filas)
                if sub is None or sub <= 0:
                    continue
                visto = True
                tasa = round((imp or 0.0) / sub, 4)
                # redondear a algo legible: 0,1300 -> 13%
                pct = round(tasa * 100, 2)
                if pct < -1 or pct > 40:
                    cta["fuera_de_rango"] += 1
                    if len(raros) < 8:
                        raros.append((os.path.basename(p)[:38], sub, imp, pct))
                    continue
                tasas[pct] += 1
        finally:
            wb.close()
        if not visto:
            cta["sin_pestana_cotizacion"] += 1

    print()
    print("=" * 62)
    print(" QUE IMPUESTO SE LE COBRO AL CLIENTE")
    print("=" * 62)
    tot = sum(tasas.values()) or 1
    print("cotizaciones con subtotal e impuesto legibles: %d" % sum(tasas.values()))
    print("  sin pestana de cotizacion: %d   ilegibles: %d   fuera de rango: %d"
          % (cta["sin_pestana_cotizacion"], cta["ilegibles"], cta["fuera_de_rango"]))
    print()
    for pct, n in sorted(tasas.items(), key=lambda x: -x[1]):
        barra = "#" * min(40, int(40.0 * n / max(tasas.values())))
        print("   %6.2f%%   %4d   %5.1f%%  %s" % (pct, n, 100.0 * n / tot, barra))
    if raros:
        print()
        print("  fuera de rango (revisar a mano):")
        for r in raros:
            print("     %-40s sub %-12s imp %-10s %.1f%%" % r)
    return 0


if __name__ == "__main__":
    sys.exit(main())

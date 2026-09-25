# -*- coding: utf-8 -*-
"""Que agrega siempre preventa aunque el cliente no lo pida.

Existe por el caso patron #5 (2026-09-25). El cliente pidio detectores,
estaciones y paneles, y las dos versiones de la cotizacion traian ademas
nueve fuentes de poder, un modulo aislador y uno de monitoreo. Nadie las
pidio porque se dan por sabidas: un sistema de deteccion no funciona sin
alimentacion. Fabian: "es parte de accesorios de montaje en todo caso".

La pestana `Compatibilidad de Accesorios` no sirve para esto: responde que
montaje SIRVE para una camara (opcional, se elige uno), no que HACE FALTA
para que un sistema funcione (obligatorio, siempre va).

En vez de escribir esas reglas de memoria, se miden. Por cada tipo de equipo
"disparador" (un NVR, un panel de incendio, una cerradura magnetica...), se
cuenta en que porcentaje de las cotizaciones que lo traen aparece tambien
cada otro tipo de linea. Si nueve de cada diez cotizaciones con NVR traen
discos duros, el disco es un accesorio requerido del NVR.

Limites honestos:
  - La clasificacion de cada linea es por palabras clave de la descripcion.
    Una linea mal descrita cae en "otro" o en el tipo equivocado.
  - Que dos cosas aparezcan juntas no prueba que una necesite a la otra:
    una camara y un NVR van juntos porque forman un sistema. Por eso el
    resultado es una lista de CANDIDATOS que alguien tiene que confirmar.

Se cuenta por COTIZACION, no por archivo: las versiones V1, V2, V3 de un
mismo proyecto se juntan en una sola.

Solo lee.

Uso:
    python scripts/medir-accesorios-requeridos.py
"""

import collections
import os
import re
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")

CLIENTES = os.path.expandvars(
    r"%USERPROFILE%\OneDrive - GRUPO VISION - DYNAMIC\Accesos directos"
    r"\Archivos de Info Costa Rica - CLIENTES"
)

# Tipo de linea -> palabras que la identifican. El orden importa: gana el
# primero que calce, asi que los mas especificos van arriba ("fuente de poder"
# antes que cualquier cosa que diga "poder").
TIPOS = [
    ("Disco duro",            r"disco\s*duro|\bhdd\b|hard\s*(disk|drive)|surveillance\s*drive|skyhawk|\bpurple\b|\bwd\d+pur"),
    ("Fuente de poder",       r"fuente\s*de\s*(poder|alimentaci)|power\s*supply|\bpsu\b|booster\s*power|fuente\s*conmutada|switching\s*power"),
    ("Bateria",               r"bater[ií]a|battery|\bbat\b"),
    ("Modulo aislador",       r"aislador|isolat"),
    ("Base de detector",      r"\bbase\b.{0,25}(detector|sensor)|(detector|sensor).{0,25}\bbase\b|\bsiga-sb\b|\bki-sb"),
    ("Panel de incendio",     r"panel.{0,30}(incendio|fire|detecci)|\bfacp\b|fire\s*alarm\s*control|\bvs4\b|\bio64|\bio1000"),
    ("Detector de incendio",  r"detector|sensor\s*de\s*(humo|temperatura|calor)|smoke|heat\s*det|\boshd\b|\bosd\b"),
    ("Estacion manual",       r"estaci[oó]n\s*manual|pull\s*station|\bm278\b|siga-278"),
    ("Sirena / estrobo",      r"estrob|strobe|sirena|horn|bocina\s*estrob|notificaci"),
    ("Cerradura magnetica",   r"cerradura|electroim[aá]n|maglock|magnetic\s*lock|\bem\s*lock"),
    ("Boton de salida",       r"bot[oó]n\s*de\s*salida|exit\s*button|request\s*to\s*exit|\brex\b|pulsador\s*de\s*salida"),
    ("Lectora",               r"lectora|lector\b|\breader\b"),
    ("Controladora de acceso", r"controlador|controller|panel\s*de\s*acceso|access\s*control\s*panel"),
    ("NVR",                   r"\bnvr\b|grabador\s*de\s*(video\s*en\s*)?red|network\s*video\s*recorder"),
    ("DVR",                   r"\bdvr\b|\bxvr\b"),
    ("Camara",                r"c[aá]mara|camera|\bbullet\b|\bdomo\b|\bdome\b|\bptz\b|turret|multisensor"),
    ("Switch",                r"\bswitch\b|conmutador"),
    ("UPS",                   r"\bups\b|no\s*break|respaldo\s*de\s*energ"),
    ("Gabinete / rack",       r"gabinete|\brack\b|cabinet"),
    ("Soporte / montaje",     r"soporte|bracket|\bmount\b|montaje|\bbrazo\b|adaptador\s*de\s*poste"),
    ("Monitor",               r"\bmonitor\b|pantalla|video\s*wall"),
    ("Licencia",              r"licencia|license"),
]
RX_TIPOS = [(t, re.compile(p, re.I)) for t, p in TIPOS]

# Los que tiene sentido mirar como "disparador": equipo principal de un
# sistema, que arrastra accesorios.
DISPARADORES = ["NVR", "DVR", "Panel de incendio", "Cerradura magnetica",
                "Controladora de acceso", "Camara", "Detector de incendio"]

PLANTILLA = re.compile(r"camara antivandalica|bodycam|face pro|lpr patrullas|"
                       r"servicio|calc mat", re.I)
NO_EQUIPO = re.compile(r"opex|mano de obra|transporte|resumen|cotiza|"
                       r"evaluacion|dise.o|tir-van", re.I)
RX_CARPETA = re.compile(r"cotizaci.n\s*#?\s*\d+\s*-\s*\d{4}", re.I)


def norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.lower()).strip()


def tipo_de(desc):
    for t, rx in RX_TIPOS:
        if rx.search(desc):
            return t
    return None


def encabezado(filas):
    for n, fila in enumerate(filas[:12]):
        cd = cc = None
        for c, v in enumerate(fila[:24]):
            t = norm(v)
            if not t or len(t) > 26:
                continue
            if cd is None and ("descripcion" in t or "description" in t):
                cd = c
            elif cc is None and ("costo unit" in t or t in ("costo u.", "costo u")):
                cc = c
        if cd is not None and cc is not None:
            return n, cd, cc
    return None


def main():
    archivos = []
    for root, dirs, files in os.walk(CLIENTES):
        if "00_IA_PREVENTAS" in root:
            continue
        for f in files:
            if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not re.search(r"matriz|oferta", f, re.I):
                continue
            p = os.path.join(root, f)
            try:
                if os.path.getsize(p) < 40000:
                    continue
            except OSError:
                continue
            archivos.append(p)
    print("matrices: %d" % len(archivos))

    por_cotizacion = collections.defaultdict(set)
    for i, p in enumerate(archivos, start=1):
        if i % 150 == 0:
            print("    %d/%d   cotizaciones con lineas: %d" % (i, len(archivos), len(por_cotizacion)))
        carpeta = os.path.dirname(p)
        d = carpeta
        while d and len(d) > len(CLIENTES):
            if RX_CARPETA.search(os.path.basename(d)):
                carpeta = d
                break
            d = os.path.dirname(d)
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            for nombre in wb.sheetnames:
                if PLANTILLA.search(nombre) or NO_EQUIPO.search(nombre):
                    continue
                try:
                    filas = list(wb[nombre].iter_rows(min_row=1, max_row=260,
                                                      values_only=True))
                except Exception:
                    continue
                cab = encabezado(filas)
                if not cab:
                    continue
                n0, cd, cc = cab
                vacias = 0
                for fila in filas[n0 + 1:]:
                    desc = fila[cd] if len(fila) > cd else None
                    costo = fila[cc] if len(fila) > cc else None
                    if not desc or not isinstance(costo, (int, float)) or costo <= 0:
                        vacias += 1
                        if vacias > 14:
                            break
                        continue
                    vacias = 0
                    t = tipo_de(str(desc))
                    if t:
                        por_cotizacion[carpeta].add(t)
        finally:
            wb.close()

    print()
    print("=" * 70)
    print(" QUE ACOMPANA A CADA EQUIPO PRINCIPAL")
    print("=" * 70)
    print("cotizaciones con al menos una linea reconocida: %d" % len(por_cotizacion))
    print()
    for disp in DISPARADORES:
        con = [s for s in por_cotizacion.values() if disp in s]
        if len(con) < 5:
            print("--- %s: solo %d cotizaciones, no alcanza ---" % (disp, len(con)))
            print()
            continue
        acomp = collections.Counter()
        for s in con:
            for t in s:
                if t != disp:
                    acomp[t] += 1
        print("--- cotizaciones con %s: %d ---" % (disp, len(con)))
        for t, n in acomp.most_common(8):
            pct = 100.0 * n / len(con)
            marca = "  <- casi siempre" if pct >= 70 else ("  <- frecuente" if pct >= 40 else "")
            print("   %-24s %4d   %5.1f%%%s" % (t, n, pct, marca))
        print()
    return 0


if __name__ == "__main__":
    sys.exit(main())

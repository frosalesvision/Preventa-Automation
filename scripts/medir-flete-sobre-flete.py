# -*- coding: utf-8 -*-
"""Cuenta cuantas veces se le aplica el 10% de transporte a una linea que YA es transporte.

Existe por el caso patron #4 (2026-09-24). Ahi un proveedor cotizo el flete
aereo como un articulo mas de la lista; esa linea entro a la matriz como
cualquier equipo y recibio encima el 10% de transporte. O sea, se le cobro
transporte al transporte.

Un caso aislado no dice nada: puede ser deliberado --un flete tambien se
nacionaliza y tambien deja margen-- o puede ser un descuido de quien copio la
lista del proveedor sin separar el flete. Esto lo cuenta sobre todas las
matrices para saber cual de las dos cosas es.

Solo lee.

Uso:
    python scripts/medir-flete-sobre-flete.py
"""

import collections
import importlib.util
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")

# El otro script tiene guiones en el nombre, asi que no se puede importar
# con `import`: hay que cargarlo por ruta. De ahi salen la carpeta de
# clientes, el detector de encabezados y el lector de lineas.
_RUTA_BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "medir-base-de-costo.py")
_spec = importlib.util.spec_from_file_location("medir_base", _RUTA_BASE)
_mb = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mb)

# Como se reconoce que una linea ES transporte. Se pide que la palabra este
# al principio o sea el grueso de la descripcion: "camara con soporte de
# transporte" no es una linea de flete.
RX_FLETE = re.compile(
    r"^\s*[-0-9]*\s*[-\s]*(transporte|flete|envio|env[ií]o|shipping|freight|courier|"
    r"carga a[eé]rea|log[ií]stica|nacionalizaci[oó]n|desalmacenaje|aduana)",
    re.I)


def columna_transporte(filas, fila_encabezado):
    """Ubica la columna de transporte por su etiqueta, arriba del encabezado."""
    for n in range(0, min(fila_encabezado, len(filas))):
        for c, v in enumerate(filas[n][:24], start=1):
            if _mb.norm(v) == "transporte":
                return c
    return None


def main():
    archivos = []
    for root, dirs, files in os.walk(_mb.CLIENTES):
        if "00_IA_PREVENTAS" in root:
            continue
        for f in files:
            if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not re.search(r"matriz|oferta", f, re.I):
                continue
            p = os.path.join(root, f)
            try:
                if os.path.getsize(p) < 40000:
                    continue
            except OSError:
                continue
            archivos.append(p)
    print("matrices a revisar: %d" % len(archivos))

    cta = collections.Counter()
    ejemplos = {"con": [], "sin": []}
    for i, p in enumerate(archivos, start=1):
        if i % 100 == 0:
            print("   %d/%d   lineas de flete halladas: %d" % (i, len(archivos), cta["flete"]))
        try:
            # read_only de verdad: este script no escribe nada nunca
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception:
            cta["ilegibles"] += 1
            continue
        try:
            for nombre in wb.sheetnames:
                try:
                    ws = wb[nombre]
                    cab = _mb.encabezado_de_equipo(ws)
                    if not cab:
                        continue
                    filas = list(ws.iter_rows(min_row=1, max_row=cab[0] + 200,
                                              values_only=True))
                    ct = columna_transporte(filas, cab[0])
                    if not ct:
                        continue
                    for n in range(cab[0], len(filas)):
                        fila = filas[n]
                        desc = fila[cab[1] - 1] if len(fila) >= cab[1] else None
                        costo = fila[cab[3] - 1] if len(fila) >= cab[3] else None
                        if not desc or not isinstance(costo, (int, float)) or costo <= 0:
                            continue
                        if not RX_FLETE.search(str(desc)):
                            continue
                        cta["flete"] += 1
                        tv = fila[ct - 1] if len(fila) >= ct else None
                        if isinstance(tv, (int, float)) and tv > 0:
                            cta["con_transporte"] += 1
                            if len(ejemplos["con"]) < 6:
                                ejemplos["con"].append((str(desc)[:44], costo, round(float(tv), 2)))
                        else:
                            cta["sin_transporte"] += 1
                            if len(ejemplos["sin"]) < 6:
                                ejemplos["sin"].append((str(desc)[:44], costo))
                except Exception:
                    continue
        finally:
            wb.close()

    print()
    print("=" * 66)
    print(" SE LE COBRA TRANSPORTE AL TRANSPORTE?")
    print("=" * 66)
    print("lineas que SON transporte/flete: %d" % cta["flete"])
    tot = cta["flete"] or 1
    print("  con el 10%% aplicado encima : %4d   %5.1f%%" % (cta["con_transporte"], 100.0*cta["con_transporte"]/tot))
    print("  sin transporte encima      : %4d   %5.1f%%" % (cta["sin_transporte"], 100.0*cta["sin_transporte"]/tot))
    print()
    for k, etq in (("con", "CON transporte encima"), ("sin", "SIN transporte encima")):
        if not ejemplos[k]:
            continue
        print("  ejemplos %s:" % etq)
        for e in ejemplos[k]:
            print("     costo %-10s %s" % (e[1], e[0]))
    return 0


if __name__ == "__main__":
    sys.exit(main())

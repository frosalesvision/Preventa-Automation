# -*- coding: utf-8 -*-
"""Revisa que el catalogo de productos este sano antes de darlo por bueno.

Existe por H-11 (2026-09-22): 29 de las 40 filas de una subcategoria no
pertenecian ahi, y nadie se dio cuenta durante semanas porque Excel no tiene
forma de avisarlo. Una fila con la subcategoria equivocada no desaparece
--sigue en el catalogo-- pero `buscar-equipo` filtra por categoria y
subcategoria, asi que deja de aparecer en las busquedas. El producto se
vuelve invisible sin que nada falle.

Correrlo DESPUES de cada carga de proveedor y DESPUES de cualquier edicion
manual del catalogo.

Uso:
    python scripts/verificar-catalogo.py
    python scripts/verificar-catalogo.py "ruta\\a\\otro\\catalogo.xlsx"

Solo lee: nunca escribe en el catalogo.
Devuelve 0 si todo esta bien, 1 si hay algo que revisar.
"""

import collections
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import taxonomia  # noqa: E402

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")


RUTA_POR_DEFECTO = os.path.expandvars(
    r"%USERPROFILE%\OneDrive - GRUPO VISION - DYNAMIC\Accesos directos"
    r"\Archivos de Info Costa Rica - CLIENTES\00_IA_PREVENTAS\Preventas"
    r"\Catalogo\Catalogo de productos por proveedor.xlsx"
)

# Columnas, base 0, segun el orden fijo documentado en actualizar-catalogo
COL_NOMBRE, COL_MARCA, COL_SKU = 0, 1, 2
COL_CATEGORIA, COL_SUBCATEGORIA, COL_PROVEEDOR = 3, 4, 5
COL_PRECIO = 6
COL_CRC, COL_DEALER_USD, COL_DEALER_CRC = 7, 8, 9


def cargar(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if "Catalogo" not in wb.sheetnames:
        sys.exit("El archivo no tiene una pestana 'Catalogo'.")
    filas = list(wb["Catalogo"].iter_rows(min_row=2, values_only=True))
    wb.close()
    return filas


def cargar_formulas(ruta):
    """Las formulas en texto, que `cargar` no ve porque pide los valores."""
    wb = openpyxl.load_workbook(ruta, read_only=True)
    hoja = wb["Catalogo"]
    out = []
    for n, f in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        out.append((n, f[COL_CRC] if len(f) > COL_CRC else None,
                    f[COL_DEALER_CRC] if len(f) > COL_DEALER_CRC else None))
    wb.close()
    return out


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else RUTA_POR_DEFECTO
    if not os.path.exists(ruta):
        sys.exit("No se encontro el catalogo en:\n  %s" % ruta)

    filas = cargar(ruta)
    validas = {cat: set(subs) for cat, subs in taxonomia.TAXONOMIA}

    print("=" * 66)
    print(" REVISION DEL CATALOGO")
    print("=" * 66)
    print("archivo: %s" % os.path.basename(ruta))
    print("filas:   %d" % len(filas))
    print()

    problemas = 0

    # --- 1. La categoria existe ---
    malas_cat = collections.Counter()
    for r in filas:
        if r[COL_CATEGORIA] not in validas:
            malas_cat[r[COL_CATEGORIA]] += 1
    print("--- 1. Categorias que no existen en la taxonomia ---")
    if malas_cat:
        for k, v in malas_cat.most_common():
            print("    %-40s %d filas" % (repr(k), v))
        problemas += sum(malas_cat.values())
    else:
        usadas = len(set(r[COL_CATEGORIA] for r in filas))
        print("    OK: las %d categorias en uso estan entre las %d de la taxonomia"
              % (usadas, len(validas)))
    print()

    # --- 2. La subcategoria corresponde a ESA categoria ---
    # Es el chequeo que importa: un desplegable plano deja poner "Montaje"
    # en una fila de camara, y eso se ve valido pero sigue estando mal.
    malos_pares = collections.Counter()
    for r in filas:
        cat, sub = r[COL_CATEGORIA], r[COL_SUBCATEGORIA]
        if cat in validas and sub not in validas[cat]:
            malos_pares[(cat, sub)] += 1
    print("--- 2. Subcategorias que no pertenecen a su categoria ---")
    if malos_pares:
        for (cat, sub), v in malos_pares.most_common():
            print("    %-30s + %-28s %d filas" % (cat, repr(sub), v))
        problemas += sum(malos_pares.values())
    else:
        print("    OK: todos los pares categoria/subcategoria son validos")
    print()

    # --- 3. Lo que el clasificador diria hoy ---
    # No es un error por si mismo: puede haber filas corregidas a mano a
    # proposito. Pero una diferencia grande avisa que el clasificador quedo
    # viejo, o que alguien reclasifico sin dejar constancia.
    difiere = []
    for i, r in enumerate(filas, 2):
        try:
            cat, sub, _ = taxonomia.clasificar(r[COL_NOMBRE], r[-1])
        except Exception:
            continue
        if (cat, sub) != (r[COL_CATEGORIA], r[COL_SUBCATEGORIA]):
            difiere.append((i, r[COL_SKU], r[COL_CATEGORIA], r[COL_SUBCATEGORIA], cat, sub))
    print("--- 3. Filas donde el clasificador diria otra cosa ---")
    if difiere:
        print("    %d filas (no es error: puede ser correccion manual a proposito)" % len(difiere))
        for i, sku, c1, s1, c2, s2 in difiere[:10]:
            print("      fila %-5d %-22s %s / %s  ->  %s / %s"
                  % (i, str(sku)[:22], c1, s1, c2, s2))
        if len(difiere) > 10:
            print("      ... y %d mas" % (len(difiere) - 10))
    else:
        print("    OK: el catalogo coincide con lo que dice el clasificador")
    print()

    # --- 4. Campos sin los que la fila no sirve para cotizar ---
    # Un SKU vacio NO es error en materiales genericos ni en servicios:
    # nadie le pone numero de parte a "tubo metal 3/4" ni a "instalacion de
    # poste". Si lo es en cualquier categoria de equipo.
    SIN_SKU_OK = {"Materiales de instalacion", "Servicios"}
    sin_sku_ok = sum(1 for r in filas
                     if not r[COL_SKU] and r[COL_CATEGORIA] in SIN_SKU_OK)
    sin_sku = sum(1 for r in filas
                  if not r[COL_SKU] and r[COL_CATEGORIA] not in SIN_SKU_OK)
    sin_precio = sum(1 for r in filas if not isinstance(r[COL_PRECIO], (int, float)))
    sin_marca = sum(1 for r in filas if not r[COL_MARCA])
    sin_prov = sum(1 for r in filas if not r[COL_PROVEEDOR])
    print("--- 4. Campos vacios que impiden cotizar esa fila ---")
    for etiqueta, n, grave in (("sin Modelo / SKU (equipo)", sin_sku, True),
                               ("sin Precio USD", sin_precio, True),
                               ("sin Proveedor", sin_prov, True),
                               ("sin Marca", sin_marca, False),
                               ("sin SKU, material o servicio", sin_sku_ok, False)):
        marca = "FALTA" if (n and grave) else ("aviso" if n else "OK   ")
        print("    %-6s %-30s %d" % (marca, etiqueta, n))
        if n and grave:
            problemas += n
    if sin_sku_ok:
        print("           (normal: materiales genericos y servicios no llevan SKU)")
    print()

    # --- 5. Mismo producto cargado dos veces ---
    #
    # Existe por un error del 2026-09-23: cargue "SS2421EM-ES" sin ver que el
    # catalogo ya tenia "SS2421EMES".  Compare las cadenas EXACTAS y un guion
    # basto para colar una fila repetida con otro precio.  Paso tres veces.
    #
    # Dos cosas que hay que tener claras para leer esto:
    #
    #   1. El SKU se compara NORMALIZADO (sin guiones, puntos ni espacios) y
    #      el nombre con los espacios colapsados.  Doce filas repetidas se
    #      escondian nada mas detras de un espacio de mas.
    #
    #   2. El catalogo es POR PROVEEDOR: el mismo producto ofrecido por dos
    #      proveedores son dos filas legitimas, y es justo lo que se quiere
    #      poder comparar al cotizar.  Un duplicado de verdad es una fila de
    #      mas DEL MISMO PROVEEDOR.
    def clave(v):
        return re.sub(r"[^A-Z0-9]", "", str(v).upper())

    def nombre(v):
        return re.sub(r"\s+", " ", str(v or "")).strip().lower()

    def precio(v):
        try:
            return round(float(v), 2)
        except (TypeError, ValueError):
            return None

    grupos = collections.defaultdict(list)
    for n, r in enumerate(filas, start=2):
        if r[COL_SKU] and clave(r[COL_SKU]):
            grupos[clave(r[COL_SKU])].append(n)

    repetidas, precio_raro, varios_prov = [], [], []
    for k, nums in grupos.items():
        if len(nums) < 2:
            continue
        por_prov = collections.defaultdict(list)
        for n in nums:
            por_prov[nombre(filas[n - 2][COL_PROVEEDOR])].append(n)
        if len(por_prov) > 1:
            varios_prov.append((k, sorted(nums)))
        for mismos in por_prov.values():
            if len(mismos) < 2:
                continue
            base = filas[mismos[0] - 2]
            for n in mismos[1:]:
                r = filas[n - 2]
                if precio(r[COL_PRECIO]) == precio(base[COL_PRECIO]):
                    repetidas.append((k, mismos[0], n, precio(base[COL_PRECIO])))
                else:
                    precio_raro.append((k, mismos[0], n,
                                        precio(base[COL_PRECIO]), precio(r[COL_PRECIO])))

    print("--- 5. Mismo producto cargado dos veces ---")
    if repetidas:
        problemas += len(repetidas)
        print("    GRAVE: %d fila(s) sobran.  Mismo proveedor, mismo SKU y el" % len(repetidas))
        print("    mismo precio: es la misma fila cargada dos veces.")
        for k, a, b, pre in repetidas[:12]:
            print("      %-22s filas %d y %d   $%s" % (k[:22], a, b, pre))
    if precio_raro:
        problemas += len(precio_raro)
        print("    REVISAR: %d caso(s) del MISMO proveedor con el mismo SKU a" % len(precio_raro))
        print("    dos precios.  O una fila quedo vieja, o son articulos")
        print("    distintos con el codigo mal puesto.  Hay que verlo a mano.")
        for k, a, b, p1, p2 in precio_raro[:12]:
            print("      %-22s filas %d y %d   $%s vs $%s" % (k[:22], a, b, p1, p2))
    if varios_prov:
        print("    %d producto(s) los ofrece mas de un proveedor.  Es normal y" % len(varios_prov))
        print("    es lo que se busca: deja comparar precio al cotizar.")
        for k, nums in varios_prov[:5]:
            print("      %-22s filas %s" % (k[:22], ", ".join(str(n) for n in nums)))
    if not (repetidas or precio_raro):
        print("    OK: ninguna fila esta cargada dos veces")
    print()

    # --- 6. Formulas de colones apuntando a otra fila ---
    #
    # Existe por lo que aparecio el 2026-09-24: 2.534 de 2.569 filas tenian
    # el precio en colones de OTRO producto.
    #
    # La causa es una trampa de openpyxl: al borrar filas NO reajusta las
    # formulas. Una que decia `G500` en la fila 500 se mueve a la 499 y
    # sigue diciendo `G500`. Cada borrado corre una fila mas el desfase, y
    # no falla nada: la celda muestra un numero perfectamente creible que
    # pertenece a otro producto.
    #
    # Por eso se revisa que cada formula mire su PROPIA fila. La de Excel
    # (COM) si reajusta bien; la trampa es solo de openpyxl.
    print("--- 6. Precio en colones apuntando a otra fila ---")
    corridas, rotas = [], []
    for n, f_crc, f_dealer in cargar_formulas(ruta):
        for col, f in (("Precio CRC", f_crc), ("Dealer CRC", f_dealer)):
            if not isinstance(f, str) or not f.startswith("="):
                continue
            if "#REF!" in f:
                rotas.append((n, col))
                continue
            m = re.search(r"IF\(([GI])(\d+)=", f)
            if m and int(m.group(2)) != n:
                corridas.append((n, col, int(m.group(2)) - n))
    if rotas:
        problemas += len(rotas)
        print("    GRAVE: %d celda(s) con #REF!.  Quedaron apuntando a una" % len(rotas))
        print("    fila que se borro.  La celda no muestra ningun precio.")
        for n, col in rotas[:6]:
            print("      fila %-6d %s" % (n, col))
    if corridas:
        problemas += len(corridas)
        print("    GRAVE: %d celda(s) muestran el precio en colones de OTRO" % len(corridas))
        print("    producto.  No da error: da un numero creible y equivocado.")
        desfases = collections.Counter(d for _, _, d in corridas)
        for d in sorted(desfases):
            print("      corridas %+d fila(s): %d celdas" % (d, desfases[d]))
        print("    Se arregla reescribiendo la formula de cada fila para que")
        print("    mire su propia fila.  El precio en dolares NO esta afectado:")
        print("    es un dato fijo, no una formula.")
    if not (rotas or corridas):
        print("    OK: cada precio en colones sale del dolar de su misma fila")
    print()

    print("=" * 66)
    if problemas == 0:
        print(" RESULTADO: el catalogo esta sano.")
    else:
        print(" RESULTADO: %d fila(s) necesitan revision.  Ver arriba." % problemas)
    print("=" * 66)
    return 1 if problemas else 0


if __name__ == "__main__":
    sys.exit(main())

# -*- coding: utf-8 -*-
"""Que marcas se cotizan de verdad y cuales no estan en el catalogo.

Existe por el caso patron #5 (2026-09-25). Ahi aparecio una alternativa en
Kidde, y el catalogo no tenia un solo producto de esa marca. La pregunta
natural fue: "no habiamos analizado que faltaba?". Si, pero por PRODUCTO, y
con un filtro que exigia que el modelo apareciera en varias cotizaciones con
precio estable. Ese filtro separo 101 modelos reales de 1.452 tokens de
ruido, pero tiene un costo: una marca que se usa en un solo proyecto, o cuyos
modelos cambian de proyecto a proyecto, se cae por diseno.

Esto mide a otra altura: por MARCA y por prefijo de modelo, contando en
cuantas cotizaciones distintas aparece cada una y cuanto se cotizo con ella.

Dos formas de reconocer la marca de una linea, porque ninguna alcanza sola:

  1. El nombre escrito en la descripcion ("Camara Hikvision ...").
  2. El prefijo del modelo ("KI-OSD" -> KI, "DS-2CD2143" -> DS). Hace falta
     porque la mayoria de las lineas NO dicen la marca: "KI-HFD - Sensor de
     temperatura" es Kidde y no lo dice en ningun lado.

Solo lee. La salida lleva nombres de proveedor y va a la carpeta temporal,
no a git.

Uso:
    python scripts/medir-marcas-faltantes.py
    python scripts/medir-marcas-faltantes.py --salida "C:\\ruta\\marcas.csv"
"""

import collections
import csv
import os
import re
import sys
import tempfile
import unicodedata
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")

CLIENTES = os.path.expandvars(
    r"%USERPROFILE%\OneDrive - GRUPO VISION - DYNAMIC\Accesos directos"
    r"\Archivos de Info Costa Rica - CLIENTES"
)
CATALOGO = os.path.join(CLIENTES, "00_IA_PREVENTAS", "Preventas", "Catalogo",
                        "Catalogo de productos por proveedor.xlsx")

# Marcas del rubro. Solo nombres publicos de fabricantes; nada de clientes ni
# precios. Se buscan como palabra entera para no confundir "HP" dentro de
# otra palabra. Se dejaron fuera nombres que son palabras comunes en una
# descripcion (Nice, Linear, Vista, Planet, Potter): daban mas ruido que
# senal.
MARCAS = {
    # video
    "Hikvision": r"hikvision|\bhik\b|hilook",
    "Dahua": r"dahua",
    "Uniview": r"uniview|\bunv\b",
    "Hanwha / Wisenet": r"hanwha|wisenet",
    "Axis": r"\baxis\b",
    "Avigilon": r"avigilon",
    "Bosch": r"\bbosch\b",
    "i-PRO / Panasonic": r"\bi-?pro\b|panasonic",
    "Vivotek": r"vivotek",
    "Milesight": r"milesight",
    "Pelco": r"\bpelco\b",
    "IDIS": r"\bidis\b",
    "InVid": r"\binvid",
    "Paramont": r"paramont",
    "Provision-ISR": r"provision",
    "Genetec": r"genetec",
    "Milestone": r"milestone",
    "ViewZ": r"viewz",
    "Epcom": r"\bepcom\b",
    "Ezviz": r"ezviz",
    # acceso
    "HID": r"\bhid\b",
    "ZKTeco": r"zkteco|\bzk\b",
    "Suprema": r"suprema",
    "Kantech": r"kantech",
    "Lenel": r"\blenel",
    "Paxton": r"paxton",
    "Assa Abloy": r"assa\s*abloy",
    "Yale": r"\byale\b",
    "Seco-Larm": r"seco-?larm|enforcer",
    "Anviz": r"anviz",
    "CDVI": r"\bcdvi\b",
    "Rosslare": r"rosslare",
    "dormakaba": r"dorma|kaba",
    # intrusion e incendio
    "DSC": r"\bdsc\b",
    "Paradox": r"paradox|digiplex",
    "Honeywell": r"honeywell|ademco",
    "Risco": r"\brisco\b",
    "Ajax": r"\bajax\b",
    "Edwards": r"edwards",
    "Simplex": r"simplex",
    "Kidde": r"kidde",
    "Notifier": r"notifier",
    "Fire-Lite": r"fire-?lite",
    "System Sensor": r"system\s*sensor",
    "Macurco": r"macurco",
    "Hochiki": r"hochiki",
    "Mircom": r"mircom",
    "Silent Knight": r"silent\s*knight",
    "Wheelock": r"wheelock",
    # redes
    "Ubiquiti / UniFi": r"ubiquiti|unifi",
    "Cisco / Meraki": r"\bcisco\b|meraki",
    "TP-Link / Omada": r"tp-?link|omada",
    "Ruckus": r"ruckus",
    "Aruba": r"\baruba\b",
    "Fortinet": r"fortinet|fortigate",
    "MikroTik": r"mikrotik",
    "Grandstream": r"grandstream",
    "Ruijie / Reyee": r"ruijie|reyee",
    "Huawei": r"huawei",
    # energia, computo, almacenamiento, cableado
    "APC": r"\bapc\b",
    "Forza": r"\bforza\b",
    "Tripp Lite": r"tripp\s*lite",
    "Eaton": r"\beaton\b",
    "CyberPower": r"cyber\s*power",
    "Western Digital": r"western\s*digital|\bwd\d",
    "Seagate": r"seagate",
    "Dell": r"\bdell\b",
    "Lenovo": r"lenovo",
    "HP": r"\bhp\b|hewlett",
    "LG": r"\blg\b",
    "Samsung": r"samsung",
    "Panduit": r"panduit",
    "Belden": r"belden",
    "CommScope": r"commscope",
    "Legrand": r"legrand",
    "Nexxt": r"\bnexxt\b",
    "Furukawa": r"furukawa",
}
RX_MARCAS = {m: re.compile(p, re.I) for m, p in MARCAS.items()}

# Prefijos que no son marca: siglas tecnicas que aparecen en cualquier
# descripcion con forma de modelo ("CAT-6", "RJ-45", "PTZ-4K"...).
NO_PREFIJO = set("""
CAT UTP FTP STP RJ PVC EMT IMC HD FHD UHD IP PTZ LED USB POE NVR DVR XVR UPS
SFP AC DC IR WDR MP TB GB MB HDMI VGA LAN WAN RS TCP WI WIFI SD CCTV ONU OLT
PDU KVA VA AWG MM CM MTS MT KG LB VCA VDC VAC IK IP66 IP67 SIM LTE GSM DSL
PC CPU RAM SSD HDD TV LCD OLED UTM RACK UL CE FCC NEMA ANSI NFPA TIPO MOD
NO SI PARA CON SIN DE LA EL LOS LAS UNA UNO POR MAS OR AND FOR THE WITH
CAM PT DOM BAL BNC RG COAX MINI MICRO MAX PRO PLUS LITE BASIC STD STANDARD
KIT SET PACK PZA UN UND PAQ ROLLO MTR ML GL LT ITEM ART COD REF SKU MODELO
""".split())

# Las pestanas de ejemplo del machote viejo traen datos de OTRO proyecto.
# Se encontraron en las siete cotizaciones revisadas el 2026-09-24.
PLANTILLA = re.compile(r"camara antivandalica|bodycam|face pro|lpr patrullas|"
                       r"servicio|calc mat", re.I)
NO_EQUIPO = re.compile(r"opex|mano de obra|transporte|resumen|cotiza|"
                       r"evaluacion|dise.o|tir-van", re.I)

RX_CARPETA = re.compile(r"cotizaci.n\s*#?\s*\d+\s*-\s*(\d{4})", re.I)
RX_TOKEN = re.compile(r"[A-Za-z0-9][A-Za-z0-9\-/\.]{3,}")


def norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.lower()).strip()


def clave(s):
    return re.sub(r"[^A-Z0-9]", "", str(s or "").upper())


def prefijo(token):
    """El prefijo de un modelo, si el token tiene forma de modelo.

    "KI-OSD" -> KI     (no tiene digitos, pero tiene guion y mayusculas)
    "DS-2CD2143G2" -> DS
    "EG1AVRF" -> EG
    "Sensor" -> None
    """
    t = token.strip(".-/")
    if len(t) < 4:
        return None
    tiene_guion = "-" in t[:7]
    tiene_digito = any(ch.isdigit() for ch in t)
    tiene_letra = any(ch.isalpha() for ch in t)
    if not tiene_letra:
        return None
    # un modelo va en mayusculas; "Camara-bullet" no es un modelo
    if t != t.upper():
        return None
    if tiene_guion:
        p = t.split("-")[0]
    elif tiene_digito:
        m = re.match(r"[A-Z]+", t)
        p = m.group(0) if m else ""
    else:
        return None
    p = p.upper()
    if len(p) < 2 or len(p) > 6 or p in NO_PREFIJO or not p.isalnum():
        return None
    if p.isdigit():
        return None
    return p


def leer_catalogo():
    wb = openpyxl.load_workbook(CATALOGO, read_only=True, data_only=True)
    filas = list(wb["Catalogo"].iter_rows(min_row=2, values_only=True))
    wb.close()
    textos = []
    pref = collections.Counter()
    for r in filas:
        textos.append(" ".join(str(x) for x in (r[0], r[1], r[13] if len(r) > 13 else "")
                               if x))
        sku = str(r[2] or "").upper()
        for tok in RX_TOKEN.findall(sku):
            p = prefijo(tok)
            if p:
                pref[p] += 1
                break
    return textos, pref


def encabezado(filas):
    """Fila de encabezado y columnas de descripcion, qty y costo unitario."""
    for n, fila in enumerate(filas[:12]):
        cd = cq = cc = None
        for c, v in enumerate(fila[:24]):
            t = norm(v)
            if not t or len(t) > 26:
                continue
            if cd is None and ("descripcion" in t or "description" in t):
                cd = c
            elif cq is None and t in ("qty", "cant", "cantidad", "cant."):
                cq = c
            elif cc is None and ("costo unit" in t or t in ("costo u.", "costo u")):
                cc = c
        if cd is not None and cc is not None:
            return n, cd, cq, cc
    return None


def main():
    salida = os.path.join(tempfile.gettempdir(), "marcas_faltantes.csv")
    if "--salida" in sys.argv:
        salida = sys.argv[sys.argv.index("--salida") + 1]

    print("leyendo el catalogo...")
    cat_textos, cat_pref = leer_catalogo()
    cat_marca = {}
    for m, rx in RX_MARCAS.items():
        cat_marca[m] = sum(1 for t in cat_textos if rx.search(t))

    print("buscando matrices...")
    archivos = []
    for root, dirs, files in os.walk(CLIENTES):
        if "00_IA_PREVENTAS" in root:
            continue
        for f in files:
            if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not re.search(r"matriz|oferta", f, re.I):
                continue
            p = os.path.join(root, f)
            try:
                if os.path.getsize(p) < 40000:
                    continue
            except OSError:
                continue
            archivos.append(p)
    print("  %d matrices" % len(archivos))

    # por (marca o prefijo) -> por carpeta de cotizacion -> monto maximo entre
    # versiones. Asi V1, V2 y V3 del mismo proyecto no se cuentan tres veces.
    monto = {"marca": collections.defaultdict(dict), "pref": collections.defaultdict(dict)}
    lineas = {"marca": collections.Counter(), "pref": collections.Counter()}
    anio = {"marca": {}, "pref": {}}
    prov = {"marca": collections.defaultdict(collections.Counter),
            "pref": collections.defaultdict(collections.Counter)}
    ejemplo = {"marca": {}, "pref": {}}
    total_lineas = 0

    for i, p in enumerate(archivos, start=1):
        if i % 150 == 0:
            print("    %d/%d   lineas leidas: %d" % (i, len(archivos), total_lineas))
        # la carpeta de la cotizacion, para no contar versiones como proyectos
        carpeta, anio_c = None, None
        d = os.path.dirname(p)
        while d and len(d) > len(CLIENTES):
            m = RX_CARPETA.search(os.path.basename(d))
            if m:
                carpeta, anio_c = d, int(m.group(1))
                break
            d = os.path.dirname(d)
        if carpeta is None:
            carpeta = os.path.dirname(p)
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            por_archivo = {"marca": collections.Counter(), "pref": collections.Counter()}
            for nombre in wb.sheetnames:
                if PLANTILLA.search(nombre) or NO_EQUIPO.search(nombre):
                    continue
                try:
                    filas = list(wb[nombre].iter_rows(min_row=1, max_row=260,
                                                      values_only=True))
                except Exception:
                    continue
                cab = encabezado(filas)
                if not cab:
                    continue
                n0, cd, cq, cc = cab
                etiqueta = ""        # la columna A marca el bloque: proveedor o seccion
                vacias = 0
                for fila in filas[n0 + 1:]:
                    if len(fila) > 0 and isinstance(fila[0], str) and fila[0].strip():
                        etiqueta = fila[0].strip()[:40]
                    desc = fila[cd] if len(fila) > cd else None
                    costo = fila[cc] if len(fila) > cc else None
                    if not desc or not isinstance(costo, (int, float)) or costo <= 0:
                        vacias += 1
                        if vacias > 14:
                            break
                        continue
                    vacias = 0
                    texto = str(desc)
                    if norm(texto).startswith(("total", "subtotal", "seccion")):
                        continue
                    qty = fila[cq] if (cq is not None and len(fila) > cq) else 1
                    qty = qty if isinstance(qty, (int, float)) and qty > 0 else 1
                    valor = float(costo) * float(qty)
                    total_lineas += 1

                    hallada = None
                    for m, rx in RX_MARCAS.items():
                        if rx.search(texto):
                            hallada = m
                            break
                    llaves = []
                    if hallada:
                        llaves.append(("marca", hallada))
                    else:
                        for tok in RX_TOKEN.findall(texto):
                            pf = prefijo(tok)
                            if pf:
                                llaves.append(("pref", pf))
                                break
                    for tipo, k in llaves:
                        lineas[tipo][k] += 1
                        por_archivo[tipo][k] += valor
                        if anio_c:
                            anio[tipo][k] = max(anio[tipo].get(k, 0), anio_c)
                        if etiqueta:
                            prov[tipo][k][etiqueta] += 1
                        if k not in ejemplo[tipo]:
                            ejemplo[tipo][k] = texto.strip()[:60]
            for tipo in ("marca", "pref"):
                for k, v in por_archivo[tipo].items():
                    previo = monto[tipo][k].get(carpeta, 0.0)
                    monto[tipo][k][carpeta] = max(previo, v)
        finally:
            wb.close()

    filas_csv = []
    for tipo in ("marca", "pref"):
        for k, porc in monto[tipo].items():
            n_cot = len(porc)
            if tipo == "pref" and n_cot < 2:
                continue        # un prefijo visto en una sola cotizacion es casi siempre ruido
            en_cat = cat_marca.get(k, 0) if tipo == "marca" else cat_pref.get(k, 0)
            top_prov = "; ".join("%s (%d)" % (e, n) for e, n in prov[tipo][k].most_common(3))
            filas_csv.append((
                "Marca" if tipo == "marca" else "Prefijo de modelo",
                k, n_cot, lineas[tipo][k], round(sum(porc.values()), 2),
                anio[tipo].get(k, ""), en_cat, top_prov, ejemplo[tipo].get(k, "")))

    filas_csv.sort(key=lambda x: (-x[2], -x[4]))
    with open(salida, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Tipo", "Marca o prefijo", "Cotizaciones", "Lineas",
                    "Monto cotizado USD (sin duplicar versiones)", "Ultimo anio",
                    "Filas en el catalogo", "Etiquetas de proveedor vistas",
                    "Ejemplo de linea"])
        w.writerows(filas_csv)

    print()
    print("=" * 74)
    print(" MARCAS QUE SE COTIZAN Y SI ESTAN EN EL CATALOGO")
    print("=" * 74)
    print("lineas de equipo leidas: %d" % total_lineas)
    print()
    for etq, tipo in (("Por MARCA escrita en la descripcion", "Marca"),
                      ("Por PREFIJO de modelo (lineas sin la marca escrita)", "Prefijo de modelo")):
        print("--- %s ---" % etq)
        print("   %-20s %5s %6s %12s %5s %8s" % ("", "cotiz", "lineas", "monto USD", "anio", "catalogo"))
        k = 0
        for f in filas_csv:
            if f[0] != tipo:
                continue
            marca_cat = "NO" if f[6] == 0 else str(f[6])
            print("   %-20s %5d %6d %12s %5s %8s" % (f[1][:20], f[2], f[3],
                  "{:,.0f}".format(f[4]), f[5], marca_cat))
            k += 1
            if k >= 30:
                break
        print()
    print("tabla completa en: %s" % salida)
    print("OJO: lleva nombres de proveedor. No subir a git.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

# -*- coding: utf-8 -*-
"""Deriva la tabla de regimenes fiscales desde las cotizaciones ya hechas.

Existe porque la hoja `Regimen Fiscal Clientes` del catalogo tiene cuatro
entradas, y al medir las 873 cotizaciones legibles resulto que el 29% van
exentas y hay 16 al 2%. O sea: hay muchas mas entidades especiales de las
que estan anotadas, y nadie se va a acordar de todas.

En vez de preguntar cliente por cliente, se saca de los documentos. Cada
cotizacion vive en la carpeta de su cliente y su pestana de cotizacion trae
el SUBTOTAL y el IMPUESTO, asi que la tasa efectiva es una division. Juntando
por cliente sale la tabla.

Un cliente puede aparecer con varias tasas y eso NO es un error: puede estar
exento en un proyecto y no en otro. Por eso se reporta la tasa dominante y
tambien las otras, para que quien revise vea la mezcla en vez de un promedio
que no significa nada.

SALIDA CON DATOS REALES DE CLIENTES. No se sube a git: se escribe donde diga
--salida, que por defecto es la carpeta temporal.

Solo lee las cotizaciones.

Uso:
    python scripts/derivar-regimen-fiscal.py
    python scripts/derivar-regimen-fiscal.py --salida "C:\\ruta\\tabla.csv"
"""

import collections
import csv
import importlib.util
import os
import re
import sys
import tempfile
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl.  Instalalo con:  pip install openpyxl")

_AQUI = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location(
    "medir_base", os.path.join(_AQUI, "medir-base-de-costo.py"))
_mb = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mb)

_spec2 = importlib.util.spec_from_file_location(
    "medir_imp", os.path.join(_AQUI, "medir-impuesto-al-cliente.py"))
_mi = importlib.util.module_from_spec(_spec2)
_spec2.loader.exec_module(_mi)

# Carpetas que no son un cliente
NO_CLIENTE = re.compile(r"^(00_|_|~|\.)", re.I)


def cliente_de(ruta):
    """El primer nivel debajo de CLIENTES es la carpeta del cliente."""
    rel = os.path.relpath(ruta, _mb.CLIENTES)
    partes = rel.split(os.sep)
    if not partes or NO_CLIENTE.match(partes[0]):
        return None
    return partes[0].strip()


def main():
    salida = os.path.join(tempfile.gettempdir(), "regimen_fiscal_derivado.csv")
    if "--salida" in sys.argv:
        salida = sys.argv[sys.argv.index("--salida") + 1]

    archivos = []
    for root, dirs, files in os.walk(_mb.CLIENTES):
        if "00_IA_PREVENTAS" in root:
            continue
        for f in files:
            if f.startswith("~$") or not f.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not re.search(r"matriz|oferta", f, re.I):
                continue
            p = os.path.join(root, f)
            try:
                if os.path.getsize(p) < 40000:
                    continue
            except OSError:
                continue
            archivos.append(p)
    print("matrices a revisar: %d" % len(archivos))

    por_cliente = collections.defaultdict(collections.Counter)
    leidas = 0
    for i, p in enumerate(archivos, start=1):
        if i % 150 == 0:
            print("   %d/%d   clientes con tasa: %d" % (i, len(archivos), len(por_cliente)))
        cli = cliente_de(p)
        if not cli:
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception:
            continue
        try:
            for nombre in wb.sheetnames:
                if not re.search(r"cotiza", nombre, re.I):
                    continue
                try:
                    filas = list(wb[nombre].iter_rows(min_row=1, max_row=90,
                                                      values_only=True))
                except Exception:
                    continue
                sub, imp = _mi.totales_de(filas)
                if sub is None or sub <= 0:
                    continue
                pct = round(((imp or 0.0) / sub) * 100, 2)
                if pct < -1 or pct > 40:
                    continue        # subtotal e impuesto en monedas distintas
                por_cliente[cli][pct] += 1
                leidas += 1
        finally:
            wb.close()

    filas_csv = []
    for cli, c in por_cliente.items():
        n = sum(c.values())
        dom, vd = c.most_common(1)[0]
        otras = "; ".join("%.2f%% x%d" % (k, v)
                          for k, v in sorted(c.items(), key=lambda x: -x[1])[1:])
        exento = "si" if dom == 0 else "no"
        # Confianza: una sola cotizacion no define un regimen.
        if n == 1:
            conf = "1 sola cotizacion - CONFIRMAR"
        elif vd == n:
            conf = "consistente en las %d" % n
        else:
            conf = "MEZCLADO: %d de %d" % (vd, n)
        filas_csv.append((cli, n, dom, exento, conf, otras))

    filas_csv.sort(key=lambda x: (-x[1], x[0]))
    with open(salida, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Cliente", "Cotizaciones", "Tasa dominante %",
                    "Exento", "Confianza", "Otras tasas vistas"])
        w.writerows(filas_csv)

    print()
    print("=" * 66)
    print(" REGIMEN FISCAL DERIVADO DE LAS COTIZACIONES")
    print("=" * 66)
    print("cotizaciones con tasa legible: %d" % leidas)
    print("clientes distintos           : %d" % len(filas_csv))
    esp = [f for f in filas_csv if f[2] != 13.0]
    print("clientes que NO van al 13%%    : %d" % len(esp))
    print()
    print("  reparto de la tasa dominante:")
    rep = collections.Counter(f[2] for f in filas_csv)
    for tasa, n in sorted(rep.items(), key=lambda x: -x[1]):
        print("     %6.2f%%   %3d clientes" % (tasa, n))
    print()
    mez = [f for f in filas_csv if f[4].startswith("MEZCLADO")]
    print("  clientes con tasas mezcladas entre proyectos: %d" % len(mez))
    print("  clientes con una sola cotizacion (a confirmar): %d"
          % len([f for f in filas_csv if f[1] == 1]))
    print()
    print("tabla escrita en: %s" % salida)
    print("OJO: lleva nombres de clientes reales. No subir a git.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

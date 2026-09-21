# -*- coding: utf-8 -*-
"""
Quita la pestana 'Proporciones Instalacion' y REHACE desde cero LEEME y
'Guia de columnas' (parchearlas en sitio rompe las celdas combinadas).

Razon de fondo (decision de Fabian, 2026-09-21): la cantidad de material
que consume una instalacion depende del sitio, no del tipo de equipo.
Guardarla como tabla le da a una estimacion la autoridad de una fuente de
verdad, y despues nadie la vuelve a cuestionar. Los materiales siguen en
el catalogo; lo que se elimina es la pretension de saber cuantos lleva.
"""
import openpyxl, warnings, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
warnings.filterwarnings("ignore")

COPIA = os.path.join(os.environ["USERPROFILE"], "Desktop", "AUTOMATIZACIONES_IA",
                     "_trabajo", "Catalogo PROPUESTA v2.xlsx")
TINTA, BANDA, SECCION, SUAVE, GRIS, BLANCO = "1F4E5F", "2E7D8F", "DCEAEE", "F4F9FA", "6B7B80", "FFFFFF"
WRAP = Alignment(wrap_text=True, vertical="top")
BORDE = Border(bottom=Side(style="thin", color="C9D8DC"))

wb = openpyxl.load_workbook(COPIA)
if "Proporciones Instalacion" in wb.sheetnames:
    del wb["Proporciones Instalacion"]
    print("pestana 'Proporciones Instalacion' eliminada")

# ======================= LEEME =======================
if "LEEME" in wb.sheetnames:
    del wb["LEEME"]
le = wb.create_sheet("LEEME", 0)
le.sheet_view.showGridLines = False
le.sheet_properties.tabColor = TINTA
le["A1"] = "Catálogo de productos por proveedor"
le["A1"].font = Font(bold=True, size=18, color=TINTA)
le.row_dimensions[1].height = 30
le["A2"] = "Todos los equipos y materiales que Grupo Visión puede cotizar, con su precio y su proveedor, en un solo lugar."
le["A2"].font = Font(size=11, italic=True, color=GRIS)
for col, w in zip("ABCD", (30, 52, 26, 30)):
    le.column_dimensions[col].width = w


def seccion(f, texto):
    le.cell(f, 1).value = texto
    le.cell(f, 1).font = Font(bold=True, size=12, color=TINTA)
    for c in range(1, 5):
        le.cell(f, c).fill = PatternFill("solid", fgColor=SECCION)
    le.row_dimensions[f].height = 22
    return f + 1


def cabecera(f, cols):
    for i, t in enumerate(cols, start=1):
        c = le.cell(f, i)
        c.value = t
        c.font = Font(bold=True, size=10, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=BANDA)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    le.row_dimensions[f].height = 26
    return f + 1


def fila(f, vals, alt):
    for i, v in enumerate(vals, start=1):
        c = le.cell(f, i)
        c.value = v
        c.font = Font(size=10.5, bold=(i == 1))
        c.alignment = WRAP
        c.border = BORDE
        if alt:
            c.fill = PatternFill("solid", fgColor=SUAVE)
    le.row_dimensions[f].height = 30
    return f + 1


f = 4
f = seccion(f, "Empezá por acá")
f = cabecera(f, ["Si querés…", "Hacé esto", "", ""])
alt = False
for p in [
    ("Buscar un equipo para cotizar", "Andá a la pestaña Catalogo y filtrá por Categoria y Subcategoria. Si no sabés qué categoría buscar, mirá el Glosario de categorias.", "", ""),
    ("Saber qué significa una columna", "Pestaña Guia de columnas: explica las columnas de cada pestaña del archivo, una por una.", "", ""),
    ("Agregar productos de un proveedor nuevo", "No los escribas a mano. Dejá el PDF o el Excel del proveedor en la carpeta Catalogos Proveedor y pedile a la IA que actualice el catálogo.", "", ""),
    ("Actualizar el tipo de cambio", "Pestaña Tipo de Cambio, celdas B3 y B4. Todos los precios en colones se recalculan solos.", "", ""),
]:
    f = fila(f, p, alt)
    alt = not alt

f += 1
f = seccion(f, "Qué hay en cada pestaña")
f = cabecera(f, ["Pestaña", "Para qué sirve", "Quién la llena", "¿Se edita a mano?"])
alt = False
for p in [
    ("LEEME", "Esta hoja. El punto de entrada para alguien que recién llega.", "Nadie, es fija", "No"),
    ("Catalogo", "El catálogo en sí: un producto por fila, con precio, proveedor y categoría.", "La IA, leyendo los documentos de los proveedores", "No. Si hay que corregir algo, pedirlo por chat"),
    ("Compatibilidad de Accesorios", "Qué montura, base o lente sirve para cada modelo de cámara.", "La IA, con un script", "No"),
    ("Tipo de Cambio", "El tipo de cambio colones/dólar que usan todas las fórmulas.", "Preventa", "SÍ — celdas B3 y B4"),
    ("Reglas de Descuento", "Qué descuento da cada proveedor y de qué depende. Determina el costo real de Grupo Visión.", "Preventa / compras", "SÍ"),
    ("Regimen Fiscal Clientes", "Qué clientes son exentos de impuesto y qué porcentaje se le cobra a cada uno.", "Preventa", "SÍ"),
    ("Tarifario Mano de Obra", "Precios fijos de instalación por unidad: por cámara, por poste, programación, capacitación.", "Preventa / operaciones", "SÍ"),
    ("Guia de columnas", "Qué significa cada columna, pestaña por pestaña.", "Nadie, es fija", "No"),
    ("Glosario de categorias", "Las categorías y subcategorías, qué incluye cada una y cómo las llama el proveedor en inglés.", "Nadie, es fija", "No"),
]:
    f = fila(f, p, alt)
    alt = not alt

f += 1
f = seccion(f, "Reglas de oro")
REGLAS = [
    "El precio del catálogo es PRECIO DE LISTA. No es lo que Grupo Visión paga. El costo real depende del descuento del proveedor y del registro del proyecto.",
    "Mirá siempre la Fecha de última actualización. Si tiene más de 2 meses, confirmá el precio con el proveedor antes de cotizar.",
    "Las columnas de precio en colones son fórmulas. No escribas un número encima: cambiá el tipo de cambio en su pestaña.",
    "No hay filas repetidas: un producto es el mismo si coinciden Modelo/SKU y Proveedor. El mismo equipo de dos proveedores son dos filas válidas.",
    "Este archivo dice QUÉ existe y CUÁNTO cuesta. No dice CUÁNTO lleva un proyecto. Los materiales de instalación (tubo, gaza, cable, cajas) están en el catálogo, pero la cantidad depende del sitio y por eso no está guardada en ningún lado: se pregunta al armar cada cotización, y se puede dejar en blanco si el proyecto no la necesita.",
]
for i, t in enumerate(REGLAS, 1):
    le.cell(f, 1).value = "%d." % i
    le.cell(f, 1).font = Font(size=10.5, bold=True)
    le.cell(f, 1).alignment = Alignment(vertical="top")
    le.cell(f, 2).value = t
    le.cell(f, 2).font = Font(size=10.5)
    le.cell(f, 2).alignment = WRAP
    le.merge_cells(start_row=f, start_column=2, end_row=f, end_column=4)
    le.row_dimensions[f].height = 30 if i < 5 else 56
    f += 1

f += 1
le.cell(f, 1).value = "¿Dudas? Preguntá por chat a la IA de preventa: puede buscar en el catálogo, explicar una columna y armar la cotización."
le.cell(f, 1).font = Font(size=11, italic=True, color=GRIS)
le.merge_cells(start_row=f, start_column=1, end_row=f, end_column=4)
print("LEEME rehecho (sin Proporciones, con la regla de oro nueva)")

# ======================= Guia de columnas =======================
CATALOGO = [
    ("Nombre de equipo/producto", "El nombre del producto tal como lo da el proveedor.", "Es lo primero que se busca y lo que se copia a la cotización.", "Siempre"),
    ("Marca", "Marca del fabricante.", "Para filtrar y para saber qué descuento aplica.", "Siempre"),
    ("Modelo / SKU", "El código exacto del proveedor.", "Con él se pide y se cotiza. Junto con Proveedor evita filas repetidas.", "Siempre"),
    ("Categoria", "Una de las 16 categorías canónicas.", "Para filtrar. El detalle está en la pestaña Glosario de categorias.", "Siempre"),
    ("Subcategoria", "El detalle dentro de la categoría (ej. Cámara → PTZ).", "Permite acotar sin leer producto por producto.", "Casi siempre"),
    ("Proveedor", "Quién nos vende ese producto.", "Define si la línea es importación y qué descuento aplica.", "Siempre"),
    ("Precio USD", "PRECIO DE LISTA en dólares. No es lo que paga Grupo Visión.", "Punto de partida del costo; el descuento se aplica al cotizar.", "Casi siempre"),
    ("Precio CRC", "El precio de lista en colones.", "FÓRMULA: se calcula con el tipo de cambio. No escribir encima.", "Automática"),
    ("Precio especial GV (USD)", "Precio de nivel distribuidor, cuando el proveedor lo publica.", "OJO: un proveedor publica tres niveles y no está confirmado cuál nos toca. No usar como costo sin preguntar.", "A veces"),
    ("Precio especial GV (CRC)", "El precio especial en colones.", "FÓRMULA: no escribir encima.", "Automática"),
    ("Unidad de venta", "Unidad, metro, rollo, caja…", "El cable se vende por metro y la cámara por unidad: cambia la cantidad a cotizar.", "Casi siempre"),
    ("Tiempo de entrega (dias)", "Cuánto tarda el proveedor en entregar.", "Para saber si se puede prometer una fecha en una licitación.", "A veces"),
    ("Pais de origen", "Dónde se fabrica.", "Para licitaciones que piden origen (ej. NDAA). NO define si es importación: eso depende del proveedor.", "A veces"),
    ("Descripcion", "El texto del proveedor, tal cual.", "Para comparar contra lo que pide el cliente.", "Casi siempre"),
    ("Especificaciones tecnicas clave", "Resolución, protección IP, certificaciones.", "Es lo que se compara contra un cartel de licitación.", "A veces"),
    ("Fecha de ultima actualizacion", "Cuándo se cargó o actualizó esta fila.", "Si tiene más de 2 meses, confirmá el precio con el proveedor.", "Siempre"),
    ("Vigencia del precio", "Hasta cuándo vale ese precio.", "Evita cotizar con un precio vencido.", "A veces"),
    ("Peso", "Peso del equipo.", "Para estimar el flete de importación.", "A veces"),
    ("Dimensiones (HxWxD)", "Medidas del equipo.", "Para el flete y para ver si cabe donde se instala.", "A veces"),
    ("Codigo HTS", "Código arancelario.", "Lo usa aduanas, no preventa.", "Casi nunca"),
    ("Codigo ECCN", "Clasificación de control de exportación.", "Lo usa logística.", "Casi nunca"),
    ("Codigo EAN", "Código de barras.", "Uso de inventario.", "Casi nunca"),
    ("Archivo de origen", "De qué documento del proveedor salió la fila.", "Para rastrear un dato que se vea raro.", "Siempre"),
    ("Pestana / Hoja de origen", "Si vino de un Excel, de qué pestaña.", "Mismo fin: rastrear el origen.", "A veces"),
    ("Categoria original del proveedor", "El texto de categoría que traía el documento original.", "SOLO REFERENCIA. Acá se ven cosas como \"IP CAMERAS & NVR's\" o \"30-Day\": así venían. Para filtrar usá Categoria.", "Siempre"),
]

BLOQUES = [
    ("Catalogo", "El catálogo en sí: un producto por fila. La llena la IA leyendo los documentos de los proveedores.", CATALOGO),
    ("Compatibilidad de Accesorios",
     "Qué montura, base o lente sirve para cada modelo de cámara. Una fila por par cámara–accesorio. La genera un script.",
     [("Modelo Camara", "El SKU de la cámara.", "Es por donde se busca: se parte del modelo que se va a cotizar.", "Siempre"),
      ("Nombre Camara", "Nombre de esa cámara.", "Para confirmar que es la que se busca.", "Siempre"),
      ("Marca Camara", "Marca de la cámara.", "Los accesorios casi nunca se comparten entre marcas.", "Siempre"),
      ("Modelo Accesorio", "El SKU del accesorio compatible.", "Es lo que se agrega a la cotización.", "Siempre"),
      ("Nombre Accesorio", "Nombre del accesorio.", "Para saber de qué tipo de montaje se trata.", "Siempre"),
      ("Tipo Accesorio", "Montura, lente, caja, etc.", "Para elegir según cómo se va a instalar (pared, poste, cielorraso).", "Siempre"),
      ("Marca Accesorio", "Marca del accesorio.", "A veces no coincide con la de la cámara y aun así sirve.", "Siempre"),
      ("Proveedor", "Quién vende el accesorio.", "Puede ser distinto al de la cámara.", "Siempre"),
      ("Fuente", "De dónde salió el par: catálogo del fabricante, match por categoría, o una cotización real.", "DICE CUÁNTO CONFIAR. Si dice categoría genérica o cotización real, verificá el encaje físico antes de cotizar.", "Siempre"),
      ("Notas", "Advertencias del par.", "Acá aparece el aviso de confirmar el encaje antes de instalar.", "A veces")]),
    ("Tipo de Cambio",
     "El tipo de cambio que usan todas las fórmulas de colones. ES DE LAS POCAS PESTAÑAS QUE SE EDITAN A MANO.",
     [("B3 — Tipo de cambio de compra", "Colones por dólar, compra.", "Se escribe a mano.", "Manual"),
      ("B4 — Tipo de cambio de venta", "Colones por dólar, venta.", "Es el que usan las fórmulas del catálogo. Si está vacío, todos los precios en colones muestran \"Actualizar TC\".", "Manual"),
      ("B5 — Fecha de actualización", "Cuándo se actualizó.", "Para saber si el tipo de cambio está viejo.", "Manual")]),
    ("Reglas de Descuento",
     "Qué descuento da cada proveedor y de qué depende. Determina el COSTO REAL de Grupo Visión. La llena preventa.",
     [("Proveedor", "A qué proveedor aplica la regla.", "Una misma marca puede tener condiciones distintas según quién la venda.", "Manual"),
      ("Marca", "A qué marca aplica.", "El descuento de una marca no es el de otra.", "Manual"),
      ("Familia / categoria", "Si la regla aplica solo a parte del catálogo.", "Vacío = aplica a toda la marca.", "Opcional"),
      ("% de descuento", "El porcentaje sobre el precio de lista.", "Es lo que convierte el precio de lista en costo real.", "Manual"),
      ("Requiere registro de proyecto", "Si el descuento depende de tener el proyecto registrado.", "Sin registro, el descuento mayor no aplica.", "Manual"),
      ("Monto minimo de compra (USD)", "Piso de compra para que aplique.", "Por debajo de ese monto el proveedor no da el descuento.", "Opcional"),
      ("Nivel de precio base", "Sobre qué precio se calcula el descuento.", "Importa cuando el proveedor publica varios niveles.", "Opcional"),
      ("Notas / Fecha / Confirmado por", "Contexto y trazabilidad.", "Para saber quién confirmó la regla y cuándo.", "Recomendado")]),
    ("Regimen Fiscal Clientes",
     "Qué impuesto se le cobra a cada cliente. OJO: no es el IVA de compra, es el que paga el cliente. La llena preventa.",
     [("Cliente", "Nombre del cliente.", "Se busca por acá al armar la cotización.", "Manual"),
      ("Exento de impuesto (si/no)", "Si al cliente no se le cobra impuesto.", "Hay instituciones exentas. Cobrar de más en una licitación es un error visible.", "Manual"),
      ("% impuesto al cliente", "Qué porcentaje se le cobra.", "No siempre es 13%: hay casos de 2%.", "Manual"),
      ("% administracion", "Si se le carga el 3% de administración.", "A unos clientes se les carga y a otros no.", "Manual"),
      ("Notas / Fecha / Confirmado por", "Contexto y trazabilidad.", "ES UN VALOR POR DEFECTO, no una verdad fija: una exención puede cambiar por contrato. Confirmarlo en cada cotización.", "Recomendado")]),
    ("Tarifario Mano de Obra",
     "Precios fijos de instalación por unidad. Solo lo que NO cambia entre proyectos: los días, las personas y la cantidad de material se preguntan en cada cotización.",
     [("Concepto", "Qué trabajo es.", "Ej. instalación de cámara, de poste, programación, capacitación.", "Manual"),
      ("Unidad", "Por qué se cobra.", "Por cámara, por poste, por día.", "Manual"),
      ("Precio USD", "Cuánto se cobra por unidad.", "Reemplaza el valor de relleno que trae la matriz.", "Manual"),
      ("Aplica a (tipo de equipo)", "A qué equipos aplica esa tarifa.", "No cuesta lo mismo instalar una domo que una PTZ en poste.", "Opcional"),
      ("Notas / Fecha / Confirmado por", "Contexto y trazabilidad.", "Para saber desde cuándo rige esa tarifa.", "Recomendado")]),
]

if "Guia de columnas" in wb.sheetnames:
    del wb["Guia de columnas"]
g = wb.create_sheet("Guia de columnas")
g.sheet_view.showGridLines = False
g["A1"] = "Guía de columnas"
g["A1"].font = Font(bold=True, size=18, color=TINTA)
g.row_dimensions[1].height = 30
g["A2"] = "Qué significa cada columna, pestaña por pestaña. Cada bloque de color oscuro es una pestaña distinta del archivo."
g["A2"].font = Font(size=11, italic=True, color=GRIS)
for col, w in zip("ABCDE", (6, 34, 46, 54, 15)):
    g.column_dimensions[col].width = w

CLAVE = {"Categoria", "Subcategoria", "Precio USD", "Precio especial GV (USD)",
         "Fecha de ultima actualizacion", "Categoria original del proveedor"}
f = 4
for pestana, para_que, cols in BLOQUES:
    g.cell(f, 1).value = "Pestaña:  %s" % pestana
    g.cell(f, 1).font = Font(bold=True, size=13, color=BLANCO)
    for c in range(1, 6):
        g.cell(f, c).fill = PatternFill("solid", fgColor=TINTA)
    g.row_dimensions[f].height = 24
    f += 1
    g.cell(f, 1).value = para_que
    g.cell(f, 1).font = Font(size=10.5, italic=True, color=TINTA)
    g.cell(f, 1).alignment = WRAP
    g.merge_cells(start_row=f, start_column=1, end_row=f, end_column=5)
    for c in range(1, 6):
        g.cell(f, c).fill = PatternFill("solid", fgColor=SECCION)
    g.row_dimensions[f].height = 28
    f += 1
    for i, t_ in enumerate(["#", "Columna", "Qué es", "Para qué sirve / qué tener en cuenta", "¿Viene llena?"], start=1):
        c = g.cell(f, i)
        c.value = t_
        c.font = Font(bold=True, size=10, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=BANDA)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    g.row_dimensions[f].height = 24
    f += 1
    for i, (nom, que, para, llena) in enumerate(cols, start=1):
        clave = pestana == "Catalogo" and nom in CLAVE
        for j, v in enumerate([i, nom, que, para, llena], start=1):
            c = g.cell(f, j)
            c.value = v
            c.font = Font(size=10.5, bold=(j == 2 and clave))
            c.alignment = WRAP
            c.border = BORDE
            if clave:
                c.fill = PatternFill("solid", fgColor=SECCION)
            elif i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=SUAVE)
        g.row_dimensions[f].height = 30
        f += 1
    f += 1

g.cell(f, 2).value = "En la pestaña Catalogo, las filas resaltadas son las columnas que más se usan al cotizar."
g.cell(f, 2).font = Font(size=11, italic=True, color=GRIS)
g.freeze_panes = "A4"
print("Guia de columnas rehecha (6 pestañas, sin Proporciones)")

orden = ["LEEME", "Catalogo", "Compatibilidad de Accesorios", "Tipo de Cambio",
         "Reglas de Descuento", "Regimen Fiscal Clientes", "Tarifario Mano de Obra",
         "Guia de columnas", "Glosario de categorias"]
wb._sheets = [wb[n] for n in orden if n in wb.sheetnames] + \
             [s for s in wb._sheets if s.title not in orden]
wb.save(COPIA)
print("")
print("pestanas finales:", [s.title for s in wb._sheets])

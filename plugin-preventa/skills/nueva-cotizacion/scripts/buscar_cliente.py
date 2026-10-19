# -*- coding: utf-8 -*-
"""
Busca un cliente en la pestana Clientes del catalogo (solo lectura) y muestra
su fila: impuesto, nota de donde salio, contacto, asesor y ultima cotizacion.
Al final muestra el tipo de cambio de la pestana Tipo de Cambio, que la ficha
propone en cada cotizacion (R7).

    python -I buscar_cliente.py "<nombre del cliente>"

Lo que imprime son datos de personas: se muestran en el chat para que el
asesor los confirme, y nunca se copian a ningun archivo del repo.
"""
import os
import sys
import unicodedata

import openpyxl

RUTA = os.path.expandvars(
    r"%USERPROFILE%\OneDrive - GRUPO VISION - DYNAMIC\Accesos directos"
    r"\Archivos de Info Costa Rica - CLIENTES\00_IA_PREVENTAS\Preventas"
    r"\Catalogo\Catalogo de productos por proveedor.xlsx"
)


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def main(buscar, ruta=RUTA):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb["Clientes"]
    filas = ws.iter_rows(values_only=True)
    enc = next(filas)
    n = 0
    for r in filas:
        if r and any(v and norm(buscar) in norm(v) for v in r[:2]):
            n += 1
            print("----")
            for i, v in enumerate(r):
                if v not in (None, ""):
                    print("  %-28s %s" % (str(enc[i] or "col %d" % (i + 1))[:28], v))
    print("%d cliente(s) con '%s'" % (n, buscar))
    print("==== Tipo de cambio (pestana Tipo de Cambio)")
    for r in wb["Tipo de Cambio"].iter_rows(values_only=True, max_row=6):
        if r and any(v not in (None, "") for v in r[:2]):
            print("  ", [v for v in r[:3] if v not in (None, "")])


if __name__ == "__main__":
    main(*sys.argv[1:3])
